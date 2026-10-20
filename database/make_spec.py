# -*- coding: utf-8 -*-
"""
database/schema.sql 기반 Oracle DB 명세서(xlsx) 생성기.
참고 규격: 여행플래너_전체화면_Oracle_DB_명세_간격조정 최종본.xlsx
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ---- 색상/서식 상수 (참고 파일과 동일) ----
NAVY   = "FF1E3A8A"   # 제목/헤더 글자·바탕
LIGHT  = "FFDBEAFE"   # 테이블명 바탕(연파랑)
GRAY   = "FFF1F5F9"   # 라벨열 바탕(연회색)
WHITE  = "FFFFFFFF"
FONT   = "NanumGothic"

def f_title():  return Font(name=FONT, size=16, bold=True,  color=NAVY)
def f_tname():  return Font(name=FONT, size=11, bold=True,  color=NAVY)
def f_head():   return Font(name=FONT, size=11, bold=True,  color=WHITE)
def f_label():  return Font(name=FONT, size=11, bold=False)
def f_data():   return Font(name=FONT, size=11, bold=False)

fill_tname = PatternFill("solid", fgColor=LIGHT)
fill_head  = PatternFill("solid", fgColor=NAVY)
fill_label = PatternFill("solid", fgColor=GRAY)

AL   = Alignment(horizontal=None, vertical="center", wrap_text=True)

ATTR_ROWS = ["의미", "Oracle 자료형", "NOT NULL (필수)", "PK / FK",
             "참조 테이블·칼럼", "기본값 / 자동증가", "추가 제약"]
# 칼럼 dict 키 순서 (의미~추가제약)
ATTR_KEYS = ["mean", "type", "notnull", "pkfk", "ref", "default", "extra"]

# ============================================================
#  테이블 정의
#  col = (칼럼명, 의미, 자료형, 필수/선택, PK/FK, 참조, 기본값, 추가제약)
# ============================================================
def C(name, mean, typ, nn, pkfk="—", ref="—", default="—", extra="—"):
    return dict(name=name, mean=mean, type=typ, notnull=nn,
                pkfk=pkfk, ref=ref, default=default, extra=extra)

D = "—"
IDENT = "IDENTITY (자동증가)"
SYST = "SYSTIMESTAMP"

MEMBER = dict(name="MEMBER", korean="회원",
  note="로그인·회원가입 인증에 사용. server/login.js 가 쓰는 칼럼과 동일하게 맞춤. 다른 테이블의 USER_ID FK 가 이 테이블을 참조.",
  tc="PRIMARY KEY (USER_ID)",
  cols=[
    C("USER_ID","아이디(로그인)","VARCHAR2(50)","필수","PK"),
    C("USER_PASSWORD","비밀번호","VARCHAR2(200)","필수"),
    C("USER_NAME","이름","VARCHAR2(50)","필수"),
    C("EMAIL","이메일","VARCHAR2(100)","선택"),
    C("CREATED_AT","가입 시간","TIMESTAMP","필수",default=SYST),
  ])

PLACE = dict(name="PLACE", korean="여행지",
  note="[슬라이드 1] 여행지 상세·비교 화면의 대표 장소(협재해변 등). 운영시간·이용요금·위치·주차 카드와 소개·지도에 사용.",
  tc="PRIMARY KEY (PLACE_ID)",
  cols=[
    C("PLACE_ID","여행지 번호","NUMBER","필수","PK",default=IDENT),
    C("NAME","이름(협재해변)","VARCHAR2(100)","필수"),
    C("REGION","지역(제주특별자치도)","VARCHAR2(50)","선택"),
    C("CITY","도시(제주시)","VARCHAR2(50)","선택"),
    C("SUMMARY","한 줄 소개","VARCHAR2(300)","선택"),
    C("DESCRIPTION","상세 소개 본문(긴 글)","CLOB","선택"),
    C("OPEN_TIME","운영시간(연중무휴 00:00~24:00)","VARCHAR2(100)","선택"),
    C("FEE","이용요금(무료)","VARCHAR2(100)","선택"),
    C("LOCATION","위치(제주시 한림읍 협재리)","VARCHAR2(200)","선택"),
    C("ADDRESS","주소","VARCHAR2(200)","선택"),
    C("PUBLIC_TRANSPORT","대중교통 안내","VARCHAR2(200)","선택"),
    C("CAR_INFO","자가용·주차 안내","VARCHAR2(200)","선택"),
    C("RECOMMEND_SEASON","추천 계절(봄·여름·가을)","VARCHAR2(100)","선택"),
    C("RECOMMEND_TIME","추천 시간대(일몰 시간)","VARCHAR2(100)","선택"),
    C("TRAVEL_STYLE","여행 스타일(가족·커플여행)","VARCHAR2(100)","선택"),
    C("KEYWORD","키워드(에메랄드빛·노을·비양도)","VARCHAR2(200)","선택"),
    C("MAP_X","지도 좌표(위도)","NUMBER(10,7)","선택"),
    C("MAP_Y","지도 좌표(경도)","NUMBER(10,7)","선택"),
    C("CREATED_AT","생성 시간","TIMESTAMP","필수",default=SYST),
  ])

PLACE_IMAGE = dict(name="PLACE_IMAGE", korean="여행지 사진",
  note="화면 상단 1/12 사진 슬라이드. SORT_ORDER 로 노출 순서를 지정.",
  tc="PRIMARY KEY (IMAGE_ID) / FOREIGN KEY (PLACE_ID) REFERENCES PLACE(PLACE_ID)",
  cols=[
    C("IMAGE_ID","사진 번호","NUMBER","필수","PK",default=IDENT),
    C("PLACE_ID","여행지 번호","NUMBER","필수","FK","PLACE.PLACE_ID"),
    C("IMAGE_URL","사진 경로","VARCHAR2(500)","필수"),
    C("SORT_ORDER","사진 순서","NUMBER","선택",default="0"),
  ])

NEARBY_PLACE = dict(name="NEARBY_PLACE", korean="주변 추천 장소",
  note="주변 추천 장소 카드(맛집·숙소·축제·행사·가볼만한 곳). CATEGORY 탭으로 필터.",
  tc="PRIMARY KEY (NEARBY_ID) / FOREIGN KEY (PLACE_ID) REFERENCES PLACE(PLACE_ID)",
  cols=[
    C("NEARBY_ID","주변 장소 번호","NUMBER","필수","PK",default=IDENT),
    C("PLACE_ID","여행지 번호","NUMBER","필수","FK","PLACE.PLACE_ID"),
    C("CATEGORY","분류(맛집/숙소/축제·행사/가볼만한 곳)","VARCHAR2(30)","선택"),
    C("NAME","이름(협재 해녀의 집)","VARCHAR2(100)","필수"),
    C("DESCRIPTION","설명(제주 향토 해산물 요리)","VARCHAR2(300)","선택"),
    C("LOCATION","위치(제주시 한림읍)","VARCHAR2(200)","선택"),
    C("IMAGE_URL","사진 경로","VARCHAR2(500)","선택"),
  ])

SAVED_PLACE = dict(name="SAVED_PLACE", korean="저장한 여행지",
  note="♡ 저장하기 / 저장한 여행지(즐겨찾기). UNIQUE 로 같은 여행지 중복 저장을 방지.",
  tc="PRIMARY KEY (SAVE_ID) / UNIQUE (USER_ID, PLACE_ID) / FK USER_ID→MEMBER, FK PLACE_ID→PLACE",
  cols=[
    C("SAVE_ID","저장 번호","NUMBER","필수","PK",default=IDENT),
    C("USER_ID","회원 아이디","VARCHAR2(50)","필수","FK","MEMBER.USER_ID"),
    C("PLACE_ID","여행지 번호","NUMBER","필수","FK","PLACE.PLACE_ID"),
    C("SAVED_AT","저장 시간","TIMESTAMP","필수",default=SYST),
  ])

ITINERARY = dict(name="ITINERARY", korean="여행 일정",
  note="AI 여행 일정 만들기 화면 왼쪽 입력 조건(기간·여행지·예산·인원·이동수단·스타일) 묶음.",
  tc="PRIMARY KEY (ITINERARY_ID) / FOREIGN KEY (USER_ID) REFERENCES MEMBER(USER_ID)",
  cols=[
    C("ITINERARY_ID","일정 번호","NUMBER","필수","PK",default=IDENT),
    C("USER_ID","회원 아이디","VARCHAR2(50)","필수","FK","MEMBER.USER_ID"),
    C("TITLE","일정 제목(2박 3일 제주 여행)","VARCHAR2(200)","선택"),
    C("START_DATE","여행 시작일","DATE","선택"),
    C("END_DATE","여행 종료일","DATE","선택"),
    C("DESTINATION","여행지(제주)","VARCHAR2(100)","선택"),
    C("BUDGET","예산(원)","NUMBER","선택"),
    C("PEOPLE_COUNT","여행 인원","NUMBER","선택"),
    C("TRANSPORT","이동 수단(자차)","VARCHAR2(50)","선택"),
    C("TRAVEL_STYLE","여행 스타일(자연·힐링)","VARCHAR2(100)","선택"),
    C("CREATED_AT","생성 시간","TIMESTAMP","필수",default=SYST),
  ])

ITINERARY_ITEM = dict(name="ITINERARY_ITEM", korean="일정 상세",
  note="1일차·2일차 시간표의 한 줄. SORT_ORDER 로 순서 바꾸기, MAP_X/Y 로 여행 경로 지도 표시. 교통·숙소 '일정에 반영' 시 이 테이블에 행이 추가됨.",
  tc="PRIMARY KEY (ITEM_ID) / FOREIGN KEY (ITINERARY_ID) REFERENCES ITINERARY(ITINERARY_ID)",
  cols=[
    C("ITEM_ID","항목 번호","NUMBER","필수","PK",default=IDENT),
    C("ITINERARY_ID","일정 번호","NUMBER","필수","FK","ITINERARY.ITINERARY_ID"),
    C("DAY_NO","일차(1/2/3일차)","NUMBER","선택"),
    C("ITEM_DATE","날짜","DATE","선택"),
    C("ITEM_TIME","시간(08:00)","VARCHAR2(10)","선택"),
    C("CATEGORY","분류(아침식사/관광/점심식사/축제·이벤트/숙소체크인)","VARCHAR2(30)","선택"),
    C("PLACE_NAME","장소명(제주 고기국수)","VARCHAR2(100)","선택"),
    C("DESCRIPTION","항목 설명","VARCHAR2(500)","선택"),
    C("LOCATION","위치(제주시 연동)","VARCHAR2(200)","선택"),
    C("DURATION","소요 시간(1시간)","VARCHAR2(50)","선택"),
    C("SORT_ORDER","순서(순서 바꾸기용)","NUMBER","선택",default="0"),
    C("MAP_X","경로 좌표(위도)","NUMBER(10,7)","선택"),
    C("MAP_Y","경로 좌표(경도)","NUMBER(10,7)","선택"),
  ])

AI_CHAT = dict(name="AI_CHAT", korean="AI 여행 도우미",
  note="오른쪽 아래 AI 여행 도우미 챗봇의 질문/답변.",
  tc="PRIMARY KEY (CHAT_ID) / FOREIGN KEY (ITINERARY_ID) REFERENCES ITINERARY(ITINERARY_ID)",
  cols=[
    C("CHAT_ID","대화 번호","NUMBER","필수","PK",default=IDENT),
    C("ITINERARY_ID","일정 번호","NUMBER","선택","FK","ITINERARY.ITINERARY_ID"),
    C("QUESTION","질문(제주도 10월 날씨는?)","VARCHAR2(1000)","선택"),
    C("ANSWER","AI 답변(긴 글)","CLOB","선택"),
    C("CREATED_AT","생성 시간","TIMESTAMP","필수",default=SYST),
  ])

AI_CHAT_SOURCE = dict(name="AI_CHAT_SOURCE", korean="참고 관광자료",
  note="AI 답변 아래의 참고 관광자료 출처 링크들.",
  tc="PRIMARY KEY (SOURCE_ID) / FOREIGN KEY (CHAT_ID) REFERENCES AI_CHAT(CHAT_ID)",
  cols=[
    C("SOURCE_ID","출처 번호","NUMBER","필수","PK",default=IDENT),
    C("CHAT_ID","대화 번호","NUMBER","필수","FK","AI_CHAT.CHAT_ID"),
    C("SOURCE_NAME","출처명(제주특별자치도 관광정보)","VARCHAR2(200)","선택"),
    C("SOURCE_URL","링크 주소","VARCHAR2(500)","선택"),
  ])

TRANSPORT = dict(name="TRANSPORT", korean="교통 정보",
  note="교통·숙소 비교 화면의 버스/기차/항공/자차 비교 목록. TAG 로 최저가·최단 시간 배지 표시.",
  tc="PRIMARY KEY (TRANSPORT_ID)",
  cols=[
    C("TRANSPORT_ID","교통편 번호","NUMBER","필수","PK",default=IDENT),
    C("TYPE","종류(버스/기차/항공/자차)","VARCHAR2(20)","선택"),
    C("COMPANY","운송사(고속버스/KTX/진에어)","VARCHAR2(100)","선택"),
    C("FROM_CITY","출발 도시(서울)","VARCHAR2(50)","선택"),
    C("TO_CITY","도착 도시(부산)","VARCHAR2(50)","선택"),
    C("DEPARTURE","출발지(서울경부터미널)","VARCHAR2(100)","선택"),
    C("ARRIVAL","도착지(부산종합터미널)","VARCHAR2(100)","선택"),
    C("DEPART_TIME","출발 시간(06:30)","VARCHAR2(10)","선택"),
    C("ARRIVE_TIME","도착 시간(10:55)","VARCHAR2(10)","선택"),
    C("DURATION","소요 시간(4시간 25분)","VARCHAR2(50)","선택"),
    C("PRICE","요금(28600)","NUMBER","선택"),
    C("TAG","태그(최저가/최단 시간)","VARCHAR2(50)","선택"),
    C("SOURCE","출처(코버스/코레일/스카이스캐너)","VARCHAR2(100)","선택"),
    C("UPDATED_AT","최종 업데이트","TIMESTAMP","필수",default=SYST),
  ])

ACCOMMODATION = dict(name="ACCOMMODATION", korean="숙소 정보",
  note="부산 인기 숙소 비교 카드. 평점·리뷰 수·1박 가격·지도 표시.",
  tc="PRIMARY KEY (ACCOM_ID)",
  cols=[
    C("ACCOM_ID","숙소 번호","NUMBER","필수","PK",default=IDENT),
    C("NAME","숙소명(부산 광안리 오션뷰 호텔)","VARCHAR2(100)","필수"),
    C("CITY","지역(부산)","VARCHAR2(50)","선택"),
    C("RATING","평점(4.7)","NUMBER(2,1)","선택"),
    C("REVIEW_COUNT","리뷰 수(832)","NUMBER","선택"),
    C("ADDRESS","주소(부산 수영구 광안해변로 219)","VARCHAR2(200)","선택"),
    C("DESCRIPTION","숙소 설명","VARCHAR2(300)","선택"),
    C("PRICE","1박 가격(120000)","NUMBER","선택"),
    C("SOURCE","출처(야놀자/아고다)","VARCHAR2(100)","선택"),
    C("MAP_X","지도 좌표(위도)","NUMBER(10,7)","선택"),
    C("MAP_Y","지도 좌표(경도)","NUMBER(10,7)","선택"),
  ])

ACCOMMODATION_AMENITY = dict(name="ACCOMMODATION_AMENITY", korean="숙소 편의시설",
  note="숙소 카드의 편의시설 태그(오션뷰·무료 Wi-Fi·조식 포함·주차 가능).",
  tc="PRIMARY KEY (AMENITY_ID) / FOREIGN KEY (ACCOM_ID) REFERENCES ACCOMMODATION(ACCOM_ID)",
  cols=[
    C("AMENITY_ID","편의시설 번호","NUMBER","필수","PK",default=IDENT),
    C("ACCOM_ID","숙소 번호","NUMBER","필수","FK","ACCOMMODATION.ACCOM_ID"),
    C("AMENITY_NAME","시설명(오션뷰/무료 Wi-Fi/조식 포함/주차 가능)","VARCHAR2(50)","선택"),
  ])

# 시트 구성: (시트명, 제목, [테이블...])
SHEETS = [
  ("여행지상세비교", "슬라이드 1 · 여행지 상세·비교 화면", [PLACE, PLACE_IMAGE, NEARBY_PLACE, SAVED_PLACE]),
  ("AI여행일정",     "슬라이드 2 · AI 여행 일정 만들기 화면", [ITINERARY, ITINERARY_ITEM, AI_CHAT, AI_CHAT_SOURCE]),
  ("교통숙소비교",   "슬라이드 3 · 교통·숙소 비교 화면", [TRANSPORT, ACCOMMODATION, ACCOMMODATION_AMENITY]),
  ("공용회원",       "공통 · 회원", [MEMBER]),
]
ALL_TABLES = [
  ("공통",     MEMBER,               "공용회원",       "로그인·회원가입 인증(기존 유지)"),
  ("슬라이드1", PLACE,               "여행지상세비교", "상세·비교 대표 장소 정보"),
  ("슬라이드1", PLACE_IMAGE,         "여행지상세비교", "상단 1/12 사진 슬라이드"),
  ("슬라이드1", NEARBY_PLACE,        "여행지상세비교", "주변 맛집·숙소·행사 카드"),
  ("슬라이드1", SAVED_PLACE,         "여행지상세비교", "♡ 저장/즐겨찾기"),
  ("슬라이드2", ITINERARY,           "AI여행일정",     "입력 조건·일정 묶음"),
  ("슬라이드2", ITINERARY_ITEM,      "AI여행일정",     "시간표 항목·여행 경로"),
  ("슬라이드2", AI_CHAT,             "AI여행일정",     "AI 도우미 질문/답변"),
  ("슬라이드2", AI_CHAT_SOURCE,      "AI여행일정",     "답변 참고자료 링크"),
  ("슬라이드3", TRANSPORT,           "교통숙소비교",   "교통편 비교 목록"),
  ("슬라이드3", ACCOMMODATION,       "교통숙소비교",   "숙소 비교 카드"),
  ("슬라이드3", ACCOMMODATION_AMENITY,"교통숙소비교",  "숙소 편의시설 태그"),
]

# 화면DB대응표 행: (화면, 이미지 요소, 연결 경로(제안), 테이블·칼럼, 표시·처리 규칙)
MAP_ROWS = [
  ("여행지상세","대표 사진 슬라이드 1/12","현재 화면","PLACE, PLACE_IMAGE","SORT_ORDER 순으로 사진 노출."),
  ("여행지상세","운영시간·이용요금·위치·주차 카드","현재 화면","PLACE.OPEN_TIME, FEE, LOCATION, CAR_INFO","상단 4개 정보 카드."),
  ("여행지상세","소개·추천(계절·시간대·스타일·키워드)","현재 화면","PLACE.SUMMARY, DESCRIPTION, RECOMMEND_SEASON, RECOMMEND_TIME, TRAVEL_STYLE, KEYWORD","소개 본문과 추천 태그."),
  ("여행지상세","위치 및 지도·대중교통·자가용","현재 화면","PLACE.ADDRESS, PUBLIC_TRANSPORT, CAR_INFO, MAP_X, MAP_Y","좌표로 지도 마커 표시."),
  ("여행지상세","주변 추천 장소(맛집/숙소/축제·행사/가볼만한 곳)","현재 화면","NEARBY_PLACE","CATEGORY 탭 필터."),
  ("여행지상세","저장하기·저장한 여행지(하트)","현재 화면","SAVED_PLACE","하트 저장/해제, 중복 저장 방지."),
  ("여행지상세","일정에 추가","/plans/new","ITINERARY, ITINERARY_ITEM","선택 여행지를 일정 항목으로 추가."),
  ("여행지상세","비슷한 여행지와 비교하기","현재 화면","PLACE","여행지 2곳의 상세정보 비교."),
  ("AI일정","왼쪽 입력(기간·여행지·예산·인원·이동수단·스타일)","현재 화면","ITINERARY","일정 생성 조건 저장."),
  ("AI일정","일정 만들기","현재 화면","ITINERARY, ITINERARY_ITEM","조건으로 일정과 항목 생성."),
  ("AI일정","1·2·3일차 시간표","현재 화면","ITINERARY_ITEM","DAY_NO·ITEM_TIME 순 정렬."),
  ("AI일정","순서 바꾸기","현재 화면","ITINERARY_ITEM.SORT_ORDER","항목 순서 변경."),
  ("AI일정","여행 경로 지도","현재 화면","ITINERARY_ITEM.MAP_X, MAP_Y","방문 순서대로 경로 표시."),
  ("AI일정","AI 여행 도우미 챗봇","현재 화면","AI_CHAT","질문/답변 저장."),
  ("AI일정","참고 관광자료 링크","현재 화면","AI_CHAT_SOURCE","답변 출처 링크 표시."),
  ("AI일정","일정 저장하기","/my-trips","ITINERARY","회원 일정으로 저장."),
  ("교통숙소","교통 비교 탭(버스/기차/항공/자차)","현재 화면","TRANSPORT","TYPE별 목록, TAG 배지 표시."),
  ("교통숙소","출발·도착·시간·요금","현재 화면","TRANSPORT.DEPARTURE, ARRIVAL, DEPART_TIME, ARRIVE_TIME, DURATION, PRICE","교통편 상세 비교."),
  ("교통숙소","숙소 비교 탭","현재 화면","ACCOMMODATION","평점·리뷰·1박 가격 카드."),
  ("교통숙소","숙소 편의시설 태그","현재 화면","ACCOMMODATION_AMENITY","오션뷰·Wi-Fi·조식·주차 태그."),
  ("교통숙소","일정에 반영","/plans/new","ITINERARY_ITEM","선택 교통/숙소를 일정 항목으로 추가."),
  ("공통","상단 메뉴·나의 여행·로그인","/login","MEMBER","회원 인증, 로그인 후 개인화."),
]

# ============================================================
#  그리기 헬퍼
# ============================================================
def set_cell(ws, coord, value, font, fill=None, align=AL):
    c = ws[coord]
    c.value = value
    c.font = font
    if fill: c.fill = fill
    c.alignment = align
    return c

def col_letter(idx):  # 1-based
    return openpyxl.utils.get_column_letter(idx)

def draw_table(ws, top, table, ncols):
    """top: 테이블명 헤더 행 번호. 15행 그리드."""
    last = col_letter(1 + ncols)       # A + 데이터칼럼수
    cols = table["cols"]

    # (1) 테이블명 행 (merge)
    ws.merge_cells(f"A{top}:{last}{top}")
    set_cell(ws, f"A{top}", f'{table["name"]} ({table["korean"]})', f_tname(), fill_tname)
    ws.row_dimensions[top].height = 24

    # (2) 칼럼명 헤더 행
    r = top + 1
    set_cell(ws, f"A{r}", "칼럼명", f_head(), fill_head)
    for j, col in enumerate(cols):
        set_cell(ws, f"{col_letter(2+j)}{r}", col["name"], f_head(), fill_head)
    ws.row_dimensions[r].height = 24

    # (3) 속성 7행
    for k, (label, key) in enumerate(zip(ATTR_ROWS, ATTR_KEYS)):
        r = top + 2 + k
        set_cell(ws, f"A{r}", label, f_label(), fill_label)
        for j, col in enumerate(cols):
            set_cell(ws, f"{col_letter(2+j)}{r}", col[key], f_data())
        ws.row_dimensions[r].height = 24

    # (4) 테이블 제약 행 (merge)
    r = top + 9
    ws.merge_cells(f"A{r}:{last}{r}")
    set_cell(ws, f"A{r}", f'테이블 제약: {table["tc"]}', f_data())
    ws.row_dimensions[r].height = 24

    # (5) 설명 행 (merge)
    r = top + 10
    ws.merge_cells(f"A{r}:{last}{r}")
    set_cell(ws, f"A{r}", table["note"], f_data())
    ws.row_dimensions[r].height = 24

def style_sheet_base(ws, ncols):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 19
    for j in range(ncols):
        ws.column_dimensions[col_letter(2+j)].width = 22

# ============================================================
#  워크북 생성
# ============================================================
wb = openpyxl.Workbook()
wb.remove(wb.active)

# ---------- 시트 1: 테이블목록 ----------
ws = wb.create_sheet("테이블목록")
ws.sheet_view.showGridLines = False
widths = [13, 33, 23, 8, 18, 46]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[col_letter(i)].width = w
ws.row_dimensions[1].height = 12
# 제목
ws.merge_cells("A2:F2")
set_cell(ws, "A2", "여행, 한국 (AI Travel Planner) · Oracle DB 명세", f_title())
ws.row_dimensions[2].height = 24
ws.row_dimensions[3].height = 12
notes = [
  "기준: 발표자료(프레젠테이션1.pptx) 3개 화면과 database/schema.sql 로 정리.",
  "전체 12개 테이블. 화면(슬라이드)별로 명세 시트를 나눔.",
  "번호 ID는 NUMBER IDENTITY 자동증가. 시간은 TIMESTAMP(기본값 SYSTIMESTAMP).",
  "회원 인증은 기존 MEMBER 테이블 기준으로 USER_ID FK 연결.",
]
for i, t in enumerate(notes):
    r = 4 + i
    ws.merge_cells(f"A{r}:F{r}")
    set_cell(ws, f"A{r}", t, f_data())
    ws.row_dimensions[r].height = 24
ws.row_dimensions[8].height = 12
# 헤더
hdr = ["구분", "테이블명", "한글 용도", "칼럼 수", "명세 시트", "요약"]
for i, h in enumerate(hdr, 1):
    set_cell(ws, f"{col_letter(i)}9", h, f_head(), fill_head)
ws.row_dimensions[9].height = 24
# 데이터
for i, (grp, tb, sheet, summary) in enumerate(ALL_TABLES):
    r = 10 + i
    set_cell(ws, f"A{r}", grp, f_data())
    set_cell(ws, f"B{r}", tb["name"], f_data())
    set_cell(ws, f"C{r}", tb["korean"], f_data())
    set_cell(ws, f"D{r}", len(tb["cols"]), f_data())
    set_cell(ws, f"E{r}", sheet, f_data())
    set_cell(ws, f"F{r}", summary, f_data())
    ws.row_dimensions[r].height = 24

# ---------- 시트 2~5: 화면별 명세 ----------
for sheet_name, title, tables in SHEETS:
    ncols = max(len(t["cols"]) for t in tables)
    ws = wb.create_sheet(sheet_name)
    style_sheet_base(ws, ncols)
    last = col_letter(1 + ncols)
    ws.row_dimensions[1].height = 12
    ws.merge_cells(f"A2:{last}2")
    set_cell(ws, "A2", f"{title} · 테이블과 칼럼", f_title())
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 12
    ws.row_dimensions[4].height = 12
    top = 5
    for tb in tables:
        draw_table(ws, top, tb, ncols)
        # 블록 사이 4행 간격
        for s in range(top + 11, top + 15):
            ws.row_dimensions[s].height = 12
        top += 15

# ---------- 시트 6: 화면DB대응표 ----------
ws = wb.create_sheet("화면DB대응표")
ws.sheet_view.showGridLines = False
mw = [10, 24, 24, 44, 50]
for i, w in enumerate(mw, 1):
    ws.column_dimensions[col_letter(i)].width = w
ws.row_dimensions[1].height = 12
ws.merge_cells("A2:E2")
set_cell(ws, "A2", "발표자료 3개 화면 · DB 대응표", f_title())
ws.row_dimensions[2].height = 24
ws.row_dimensions[3].height = 12
ws.merge_cells("A4:E4")
set_cell(ws, "A4", "데이터는 아래 테이블로 연결. 이동 버튼·지도 조작·예시 이미지는 라우팅·프론트 구현. 화면 속 사진·날짜·가격·순위는 예시 표기.", f_data())
ws.row_dimensions[4].height = 24
ws.row_dimensions[5].height = 12
mhdr = ["화면", "이미지 요소", "연결 경로(제안)", "테이블·칼럼", "표시·처리 규칙"]
for i, h in enumerate(mhdr, 1):
    set_cell(ws, f"{col_letter(i)}6", h, f_head(), fill_head)
ws.row_dimensions[6].height = 24
for i, row in enumerate(MAP_ROWS):
    r = 7 + i
    for j, val in enumerate(row, 1):
        set_cell(ws, f"{col_letter(j)}{r}", val, f_data())
    ws.row_dimensions[r].height = 24

# ---------- 저장 ----------
import os
out = os.path.join(os.path.dirname(__file__), "여행한국_Oracle_DB_명세.xlsx")
wb.save(out)
print("saved:", out)
print("sheets:", wb.sheetnames)
